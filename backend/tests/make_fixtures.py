"""
Generates the sample documents used by the tests and for manual checks.

Run from the `backend/` directory:

    .venv/Scripts/python.exe tests/make_fixtures.py

Creates files in tests/fixtures/:
    sample.txt          plain text, multiple paragraphs
    sample.docx         headings, paragraphs and a table
    sample.pdf          real text layer, 3 pages
    scanned.pdf         image-only pages, no text layer (for OCR in Phase 3)
    empty.txt           whitespace only (for the empty-document test)
"""

from pathlib import Path

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from PIL import Image, ImageDraw

FIXTURES = Path(__file__).parent / "fixtures"

PARAGRAPHS = [
    "Artificial intelligence has transformed the way organizations process "
    "unstructured information. Large volumes of text are produced every day in "
    "the form of reports, research papers and legal contracts. Reading all of "
    "this material manually is slow and expensive.",
    "Automatic text summarization addresses this problem by producing a shorter "
    "version of a document that preserves its most important information. "
    "Extractive methods select entire sentences from the source text, while "
    "abstractive methods generate new sentences that may use words not present "
    "in the original.",
    "Extractive summarization is reliable because every sentence it returns is "
    "guaranteed to appear in the source document. It cannot introduce "
    "information that is not there. Its weakness is that the selected sentences "
    "may not read fluently together.",
    "Abstractive summarization produces more natural and compact text. However, "
    "neural models can generate statements that are not supported by the source "
    "document, a problem known as hallucination. Hybrid systems reduce this risk "
    "by selecting important sentences first and rewriting only that material.",
    "Evaluation of summaries is usually performed with ROUGE and BLEU. Both "
    "metrics compare a generated summary against a reference summary written by "
    "a human. They therefore cannot be used when no reference summary exists.",
]


def write_txt() -> None:
    (FIXTURES / "sample.txt").write_text("\n\n".join(PARAGRAPHS), encoding="utf-8")


def write_empty_txt() -> None:
    (FIXTURES / "empty.txt").write_text("   \n\n \t \n", encoding="utf-8")


def write_docx() -> None:
    document = DocxDocument()
    document.add_heading("Automatic Text Summarization", level=1)
    for paragraph in PARAGRAPHS[:3]:
        document.add_paragraph(paragraph)

    document.add_heading("Evaluation metrics", level=2)
    table = document.add_table(rows=3, cols=2)
    rows = [
        ("Metric", "What it measures"),
        ("ROUGE", "Overlap with a reference summary"),
        ("BLEU", "Precision of matching n-grams"),
    ]
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value

    document.save(FIXTURES / "sample.docx")


def write_pdf() -> None:
    """A normal PDF with a real text layer."""
    document = fitz.open()
    for paragraph in PARAGRAPHS[:3]:
        page = document.new_page()
        page.insert_textbox(fitz.Rect(50, 50, 545, 780), paragraph, fontsize=11)
    document.save(FIXTURES / "sample.pdf")
    document.close()


def write_scanned_pdf() -> None:
    """A PDF whose pages are images only - no text layer at all.

    Used to verify that a scanned document is detected, and (from Phase 3) that
    OCR recovers the text.
    """
    scanned_text = [
        "This page is an image. Optical character recognition is required to "
        "read it. Tesseract will convert the pixels into characters.",
        "The second scanned page tests that multiple pages are processed and "
        "combined in the correct order.",
    ]

    document = fitz.open()
    for content in scanned_text:
        image = Image.new("RGB", (1200, 400), "white")
        draw = ImageDraw.Draw(image)
        # Wrap manually: PIL's default bitmap font has fixed metrics.
        lines, line = [], ""
        for word in content.split():
            if len(line) + len(word) > 62:
                lines.append(line)
                line = word
            else:
                line = f"{line} {word}".strip()
        lines.append(line)
        for index, text in enumerate(lines):
            draw.text((30, 40 + index * 26), text, fill="black")

        image_path = FIXTURES / "_scanned_page.png"
        image.save(image_path)

        page = document.new_page(width=595, height=300)
        page.insert_image(page.rect, filename=str(image_path))

    document.save(FIXTURES / "scanned.pdf")
    document.close()
    (FIXTURES / "_scanned_page.png").unlink(missing_ok=True)


def main() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    write_txt()
    write_empty_txt()
    write_docx()
    write_pdf()
    write_scanned_pdf()

    for path in sorted(FIXTURES.iterdir()):
        print(f"  {path.name:15} {path.stat().st_size:>8,} bytes")


if __name__ == "__main__":
    main()
